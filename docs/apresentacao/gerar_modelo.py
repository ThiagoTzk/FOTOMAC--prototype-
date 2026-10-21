"""Builds the FOTOMAC presentation on top of CESMAC's official template.

    ~/Agente/skills/apresentacoes/.venv/bin/python docs/apresentacao/gerar_modelo.py

Input:  docs/entregaveis/MODELO - APRESENTAÇÃO REGISTRO.pptx (4:3, CESMAC logo and colour band
        live in the template's master, so every slide gets them for free).
Output: docs/entregaveis/FOTOMAC-apresentacao-modelo.pptx

The template's own slides are edited in place (title boxes, bullet lists and fonts are kept as
they are: Segoe UI Semibold titles, Segoe UI bullets, Arial on the cover and the software sheet).
New slides clone the template's title and bullet boxes instead of imitating them, so they match
exactly. The blue boxes of the template are instructions to the student and are removed.
Every number here comes from the code, the requirements v1.6 or fotomac-lab's Test 1.
"""
from __future__ import annotations

import copy

from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
DOCS = HERE.parent
TEMPLATE = DOCS / "entregaveis" / "MODELO - APRESENTAÇÃO REGISTRO.pptx"
OUT = DOCS / "entregaveis" / "FOTOMAC-apresentacao-modelo.pptx"
SHOT_SITE = DOCS / "entregaveis" / "Captura de tela Representação visual Insight face (explicação breve e visual do insight).png"
SHOT_CONTACT = DOCS / "entregaveis" / "Captura de tela contato corporativo (representação de como fazer o contato com a empresa).png"

# Colours: CESMAC/FOTOMAC blues and yellow (the template band) + the template's purple bullet.
BLUE, NAVY, YELLOW = "007DC5", "072D44", "FFC506"
INK, MUTED, LINE, PAGE, SOFT = "1E2933", "65727E", "DDE3E8", "F5F7F9", "E8F5FC"
GREEN, GREEN_SOFT, AMBER, AMBER_SOFT = "23855B", "E8F7F0", "9A6500", "FFF4D6"
PURPLE, PURPLE_SOFT = "6D28D9", "F1ECFD"
FONT = "Arial"

# Content must stay between the title and the colour band of the master (band starts at ~6.84").
LEFT, RIGHT = 0.5, 9.5
WIDTH = RIGHT - LEFT

_FONTS = "/System/Library/Fonts/Supplemental/"
_metric = {False: ImageFont.truetype(_FONTS + "Arial.ttf", 100), True: ImageFont.truetype(_FONTS + "Arial Bold.ttf", 100)}


def text_width_in(text: str, size_pt: float, bold: bool = False, segoe: bool = False) -> float:
    """Approximate width in inches. Segoe UI is not on this Mac: Arial metrics + 6% margin."""
    w = _metric[bold].getlength(text) / 100 * size_pt / 72
    return w * (1.06 if segoe else 1.0)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# ------------------------------------------------------------------ template plumbing
prs = Presentation(str(TEMPLATE))
T = list(prs.slides)  # T[0] = template slide 1 ... T[9] = slide 10


def shape_named(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def remove(shape):
    shape._element.getparent().remove(shape._element)


TITLE_PROTO = copy.deepcopy(shape_named(T[2], "TextShape 2")._element)   # "INTRODUÇÃO"
BULLET_PROTO = copy.deepcopy(shape_named(T[2], "CustomShape 1")._element)  # purple bullets
BLANK_LAYOUT = T[2].slide_layout


def set_title(slide, text, shape=None, top=None):
    """Write a title keeping the template's run formatting; shrink only if it would not fit."""
    if shape is None:
        el = copy.deepcopy(TITLE_PROTO)
        slide.shapes._spTree.append(el)
        shape = next(s for s in slide.shapes if s._element is el)
    if top is not None:
        shape.top = Inches(top)
    run = shape.text_frame.paragraphs[0].runs[0]
    run.text = text
    size = 44
    while text_width_in(text, size, bold=True, segoe=True) > WIDTH - 0.3 and size > 30:
        size -= 2
    run.font.size = Pt(size)
    return shape


def set_bullets(shape, items, size=None):
    """Replace the paragraphs of a template bullet box, cloning its first paragraph's format."""
    txBody = shape.text_frame._txBody
    paras = txBody.findall(qn("a:p"))
    proto = copy.deepcopy(paras[0])
    for p in paras:
        txBody.remove(p)
    for item in items:
        p = copy.deepcopy(proto)
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        for e in p.findall(qn("a:endParaRPr")):
            p.remove(e)
        runs[0].find(qn("a:t")).text = item
        if size:
            runs[0].find(qn("a:rPr")).set("sz", str(int(size * 100)))
        txBody.append(p)


def new_slide(title, title_top=0.8):
    slide = prs.slides.add_slide(BLANK_LAYOUT)
    for ph in list(slide.placeholders):
        remove(ph)
    set_title(slide, title, top=title_top)
    return slide


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def tx(slide, x, y, w, h, runs, size=14, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, spacing=0):
    """Text box. `runs` is a string, or a list of paragraphs; each paragraph a string or a list of
    (text, {font overrides}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing and i:
            p.space_before = Pt(spacing)
        for text, over in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = text
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(over.get("size", size)), over.get("bold", bold), over.get("italic", italic)
            f.color.rgb = rgb(over.get("color", color))
    return box


def box(slide, x, y, w, h, fill, line=None, radius=0.12, dashed=False, line_w=1.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = min(0.5, radius / min(w, h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dashed:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def circle(slide, x, y, d, fill, label, size=14):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
    r.font.color.rgb = rgb("FFFFFF")


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def labelled_box(slide, x, y, w, h, title, sub, fill, color, dashed=False, title_size=16, sub_size=13):
    box(slide, x, y, w, h, fill, color, dashed=dashed)
    tx(slide, x + 0.1, y, w - 0.2, h, [[(title, {"bold": True, "size": title_size, "color": color})],
                                       [(sub, {"size": sub_size, "color": INK})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=2)


# ================================================================== 1. cover (template slide 1)
cover = T[0]
title_box = shape_named(cover, "TextShape 2")
runs = title_box.text_frame.paragraphs[0].runs
runs[0].text = "FOTOMAC: "
runs[1].text = "ACERVO FOTOGRÁFICO COM BUSCA POR RECONHECIMENTO FACIAL"
for r in runs:
    r.font.size = Pt(30)

people = shape_named(cover, "CustomShape 1")
people.left, people.width = Inches(0.4), Inches(9.2)
credit = [("AUTOR", "Thiago José Camêlo Nunes"),
          ("COLABORADOR", "Pedro Rosa"),
          ("ORIENTADOR", "Prof. Me. Wagner de Oliveira Lima Palmeira de Araujo"),
          ("STAKEHOLDERS", "Prof. Me. Wagner, Moacir Barbosa (fotógrafo) e CESMAC")]
txBody = people.text_frame._txBody
paras = txBody.findall(qn("a:p"))
proto = copy.deepcopy(paras[0])  # bold label run + regular value run, Arial 18, centred
for p in paras:
    txBody.remove(p)
for label, value in credit:
    p = copy.deepcopy(proto)
    r_label, r_value = p.findall(qn("a:r"))[:2]
    r_label.find(qn("a:t")).text = label
    r_value.find(qn("a:t")).text = f": {value}"
    for r in (r_label, r_value):
        r.find(qn("a:rPr")).set("sz", "1600")
    txBody.append(p)
notes(cover, "Apresentar o FOTOMAC em uma frase: a pessoa encontra as próprias fotos dos eventos do CESMAC pelo "
             "rosto, sem ver as fotos dos outros. Agradecer ao Pedro Rosa, que ajudou em boa parte do desenvolvimento, "
             "e citar os clientes: o fotógrafo Moacir Barbosa e o CESMAC.")

# ================================================================== 2. roteiro (template slide 2)
roteiro = T[1]
remove(shape_named(roteiro, "CaixaDeTexto 1"))
set_bullets(shape_named(roteiro, "CustomShape 1"),
            ["Introdução", "Informações do Software", "Metodologia e Desenvolvimento", "Vídeo",
             "Resultados Parciais", "Referências"])
notes(roteiro, "Roteiro rápido: o problema e a solução, a ficha do software, como foi construído, o vídeo do "
               "sistema funcionando e o que está pronto e o que falta.")

# ================================================================== 3. introdução (template slide 3)
intro = T[2]
for name in ("CaixaDeTexto 3", "CaixaDeTexto 4"):
    remove(shape_named(intro, name))
body = shape_named(intro, "CustomShape 1")
body.top = Inches(2.0)
set_bullets(body, ["O problema: folhear todas as fotos",
                   "Até 1.000 fotos por evento",
                   "Quem procura vê o rosto de todos",
                   "Solução: uma selfie, só as suas fotos",
                   "Rosto é dado sensível (LGPD)"])
notes(intro, "Um evento do CESMAC pode ter até mil fotos. Hoje, para achar as suas, você olha todas — e quem "
             "procura vê o rosto de todo mundo. No FOTOMAC, o visitante lê o aviso, tira uma selfie e recebe só as "
             "fotos em que aparece. Como rosto é dado pessoal sensível pela LGPD (Art. 5º, II), a privacidade entrou "
             "no desenho desde o começo.")

# ================================================================== 4. informações do software (template slide 4)
info = T[3]
sheet = shape_named(info, "CustomShape 1")
values = {
    "Nome do Produto:": "FOTOMAC (denominação provisória)",
    "Linguagem de Programação:": "Python (servidor) e TypeScript (site)",
    "Banco de dados:": "PostgreSQL com pgvector; as fotos ficam no SeaweedFS",
    "Resumo da Aplicabilidade do Software:": "acervo fotográfico dos eventos do CESMAC em que a própria pessoa "
                                             "encontra as suas fotos por reconhecimento facial",
    "Qual problema o software pretende resolver:": "achar as próprias fotos entre até 1.000 por evento sem expor "
                                                   "o rosto de todos",
    "Framework:": "FastAPI (servidor), Expo / React Native (site; o aplicativo vem depois) e InsightFace "
                  "(reconhecimento facial)",
}
for p in sheet.text_frame.paragraphs:
    label = "".join(r.text for r in p.runs).strip()
    key = next((k for k in values if label.startswith(k)), None)
    if key is None:
        continue
    for r in p.runs[1:]:  # drops the "(caso haja)" hint
        r._r.getparent().remove(r._r)
    p.runs[0].text = key
    value = copy.deepcopy(p.runs[0]._r)
    value.find(qn("a:t")).text = " " + values[key]
    value.find(qn("a:rPr")).set("b", "0")
    p.runs[0]._r.addnext(value)
notes(info, "A ficha do FOTOMAC: servidor em Python com FastAPI, site em TypeScript com Expo — a mesma base vira o "
            "aplicativo depois. Banco PostgreSQL com a extensão pgvector, que guarda e compara os vetores de rosto. "
            "As fotos ficam no SeaweedFS, o armazenamento que o servidor do CESMAC já usa.")

# ================================================================== 5. metodologia (template slide 5)
metodo = T[4]
# 44 pt only fits in the template because the preview swaps Segoe UI for the narrower Times.
set_title(metodo, "Metodologia e Desenvolvimento", shape=shape_named(metodo, "TextShape 1"))
set_bullets(shape_named(metodo, "CustomShape 1"),
            ["Requisitos com o cliente (v1.6)",
             "Construção em etapas, cada uma testada",
             "IA escolhida por teste com fotos reais",
             "LGPD aplicada desde o desenho"])
notes(metodo, "Primeiro levantei os requisitos com o cliente — o documento está na versão 1.6. Depois construí em "
              "etapas, cada uma com testes automáticos. A IA de reconhecimento facial foi escolhida por um teste meu, "
              "com fotos reais de formatura. E a LGPD entrou em cada decisão, não como revisão no fim.")

# ================================================================== 6. arquitetura (new)
arq = new_slide("Arquitetura")
labelled_box(arq, 3.0, 1.95, 4.0, 0.8, "Site · Expo", "o que a pessoa vê no navegador", PAGE, "334155")
labelled_box(arq, 0.9, 3.2, 8.2, 0.85, "API · Python + FastAPI", "recebe todo pedido, confere a permissão e decide",
             SOFT, BLUE)
labelled_box(arq, 0.5, 4.6, 2.85, 1.05, "PostgreSQL + pgvector", "fichas e vetores de rosto", GREEN_SOFT, GREEN,
             title_size=15)
labelled_box(arq, 3.575, 4.6, 2.85, 1.05, "InsightFace", "reconhece os rostos", PURPLE_SOFT, PURPLE, dashed=True)
labelled_box(arq, 6.65, 4.6, 2.85, 1.05, "SeaweedFS", "guarda as fotos (S3)", AMBER_SOFT, AMBER)
arrow(arq, 5.0, 2.75, 5.0, 3.17)
tx(arq, 5.15, 2.8, 2.0, 0.3, "HTTP · JSON", size=11, color=MUTED)
arrow(arq, 2.2, 4.05, 1.93, 4.57)
arrow(arq, 5.0, 4.05, 5.0, 4.57)
arrow(arq, 7.8, 4.05, 8.07, 4.57)
tx(arq, 2.5, 5.8, 5.0, 0.35, "adaptador: troca a IA sem mexer no resto", size=13, color=PURPLE, italic=True,
   align=PP_ALIGN.CENTER)
notes(arq, "O site só conversa com a API. A API fala com o banco, com o reconhecimento facial e com o armazenamento "
           "das fotos — e confere a permissão antes. O InsightFace fica atrás de um adaptador: trocar de IA é "
           "escrever um adaptador novo, sem mexer no resto.")

# ================================================================== 7. reconhecimento facial (new)
rec = new_slide("Reconhecimento facial")
steps = [("Detecta", "SCRFD"), ("Alinha", "olhos no lugar"), ("Gera o vetor", "ArcFace · 512 números"),
         ("Compara", "similaridade")]
w, gap, y = 2.0, 0.333, 2.05
for i, (t, sub) in enumerate(steps):
    x = LEFT + i * (w + gap)
    last = i == 3
    labelled_box(rec, x, y, w, 1.15, t, sub, SOFT if last else PURPLE_SOFT, BLUE if last else PURPLE,
                 title_size=16, sub_size=12)
    if not last:
        arrow(rec, x + w + 0.04, y + 0.575, x + w + gap - 0.04, y + 0.575)
tx(rec, LEFT, 3.75, 2.4, 1.1, "0,30", size=60, bold=True, color=BLUE)
tx(rec, 3.0, 3.9, 6.5, 1.0, [[("limiar de similaridade", {"bold": True, "color": INK})],
                             "abaixo dele, a foto não aparece:",
                             "melhor não mostrar do que mostrar a foto errada"], size=15, color=MUTED, spacing=2)
box(rec, 6.9, 5.35, 2.6, 0.5, PURPLE_SOFT, radius=0.25)
tx(rec, 6.9, 5.35, 2.6, 0.5, "InsightFace buffalo_l", size=13, bold=True, color=PURPLE, align=PP_ALIGN.CENTER,
   anchor=MSO_ANCHOR.MIDDLE)
notes(rec, "O InsightFace acha o rosto (SCRFD), alinha e transforma em 512 números (ArcFace). Para comparar dois "
           "rostos, mede-se o quanto esses vetores apontam para o mesmo lado. Só aparece o que passa de 0,30 — no "
           "Teste 1, a foto errada mais alta tirou 0,278.")

# ================================================================== 8. Test 1 (new)
t1 = new_slide("Escolha da IA: Teste 1")
tx(t1, LEFT, 1.85, 6.0, 0.3, "Nota no Teste 1 (0 a 100)", size=12, color=MUTED)
rows = [("InsightFace", 97.6), ("YuNet + SFace", 89.5), ("ArcFace", 64.0), ("Facenet512", 60.3), ("VGG-Face", 57.1)]
bx, max_w = 2.35, 3.6
for i, (name, v) in enumerate(rows):
    y, top = 2.3 + i * 0.66, i == 0
    tx(t1, LEFT, y, 1.7, 0.46, name, size=14, bold=top, color=INK, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    box(t1, bx, y, max_w * v / 100, 0.46, BLUE if top else "C9D3DB", radius=0.05)
    tx(t1, bx + max_w * v / 100 + 0.1, y, 0.8, 0.46, f"{v:.1f}".replace(".", ","), size=14, bold=True,
       color=BLUE if top else INK, anchor=MSO_ANCHOR.MIDDLE)
for y, big, small in [(2.2, "388", "fotos reais conferidas à mão"), (3.65, "7 × 203", "fotos erradas no topo: 1º × 2º")]:
    tx(t1, 6.9, y, 2.6, 0.7, big, size=36, bold=True, color=BLUE)
    tx(t1, 6.9, y + 0.72, 2.6, 0.5, small, size=13, color=MUTED)
tx(t1, 6.9, 5.3, 2.6, 0.5, "5 tecnologias · fotos de formatura · fotomac-lab", size=11, italic=True, color=MUTED)
notes(t1, "Antes do sistema, montei uma bancada e comparei 5 tecnologias em 388 fotos reais de formatura, conferindo "
          "cada uma à mão. O InsightFace venceu com folga: 7 fotos erradas no topo do ranking contra 203 do segundo. "
          "Isso contraria as tabelas publicadas — elas medem rostos recortados, não plateia.")

# ================================================================== 9. why not generative AI (new)
gen = new_slide("Por que não uma IA generativa?")
cw, ax, bxx = 3.45, 2.5, 6.05
for x, t, sub, fill, color in [(ax, "IA generativa", "ChatGPT · Gemini · Llama Vision", AMBER_SOFT, AMBER),
                               (bxx, "InsightFace", "IA especializada em rosto", BLUE, "FFFFFF")]:
    box(gen, x, 1.95, cw, 0.7, fill, radius=0.08)
    tx(gen, x, 1.95, cw, 0.7, [[(t, {"bold": True, "size": 16, "color": color})], [(sub, {"size": 11, "color": color})]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
table = [("Máquina", "GPU com 11 a 22 GB de memória¹", "Só processador · ~1 GB de RAM"),
         ("Resposta", "Texto livre, que varia e pode errar", "512 números, sempre os mesmos"),
         ("Busca em 1.000 fotos", "Rever as fotos a cada busca", "Calcula uma vez, ao publicar"),
         ("Privacidade", "Na nuvem, o rosto vai para outra empresa, fora do país", "O rosto não sai do servidor")]
for i, (label, a, b) in enumerate(table):
    y = 2.8 + i * 0.85
    line = gen.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(LEFT), Inches(y + 0.8), Inches(RIGHT), Inches(y + 0.8))
    line.line.color.rgb, line.line.width = rgb(LINE), Pt(0.75)
    tx(gen, LEFT, y, 1.9, 0.78, label, size=14, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    tx(gen, ax + 0.1, y, cw - 0.2, 0.78, a, size=14, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tx(gen, bxx + 0.1, y, cw - 0.2, 0.78, b, size=14, bold=True, color=BLUE, align=PP_ALIGN.CENTER,
       anchor=MSO_ANCHOR.MIDDLE)
tx(gen, LEFT, 6.25, WIDTH, 0.45, "¹ Llama 3.2 Vision 11B rodando local (estimativa WILL IT RUN AI, 2026). InsightFace medido "
                                 "neste projeto: 0,13 s por foto, notebook sem GPU.", size=10, italic=True, color=MUTED)
notes(gen, "O InsightFace também é IA — mas especializada: só faz rosto. Uma IA generativa, como o ChatGPT, foi feita "
           "para conversar. Para rodar uma dessas aqui dentro, precisaria de placa de vídeo com 11 a 22 GB; o "
           "InsightFace roda no processador com cerca de 1 GB. Ela responde em texto, que muda de uma vez para outra; o "
           "InsightFace devolve sempre os mesmos 512 números, calculados uma vez quando a foto é publicada. E usar uma "
           "IA na nuvem seria mandar rosto, dado sensível, para outra empresa fora do Brasil — a LGPD só permite "
           "transferência internacional nos casos do Art. 33.")

# ================================================================== 10. licence (new)
lic = new_slide("Licença do InsightFace")
tx(lic, LEFT, 1.85, WIDTH, 0.4, "Livre para pesquisar, licenciado para operar", size=20, bold=True, color=BLUE,
   align=PP_ALIGN.CENTER)
lic.shapes.add_picture(str(SHOT_SITE), Inches(LEFT), Inches(2.45), width=Inches(4.6))
tx(lic, LEFT, 5.13, 4.6, 0.3, "insightface.ai: a empresa vende licença para uso comercial", size=10, italic=True,
   color=MUTED)
for i, (t, sub, color) in enumerate([("Código: MIT", "livre, inclusive para uso comercial", GREEN),
                                     ("Modelos treinados: só pesquisa", "o buffalo_l deste TCC está coberto", GREEN),
                                     ("Operar no CESMAC: exige licença", "decisão da instituição, não do TCC", AMBER)]):
    y = 2.5 + i * 0.9
    box(lic, 5.35, y, 0.07, 0.68, color, radius=0.03)
    tx(lic, 5.55, y, 3.95, 0.7, [[(t, {"bold": True, "size": 15, "color": NAVY})],
                                 [(sub, {"size": 13, "color": MUTED})]], spacing=2)
box(lic, LEFT, 5.55, WIDTH, 0.8, SOFT, radius=0.12)
tx(lic, LEFT + 0.2, 5.55, WIDTH - 0.4, 0.8, "Plano B já medido: YuNet + SFace (licenças MIT e Apache 2.0, uso comercial "
                                            "livre), nota 89,5 no Teste 1. O adaptador troca o modelo sem reescrever o "
                                            "sistema.", size=13, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(lic, "O código do InsightFace é MIT, sem restrição. O que tem restrição são os modelos treinados, como o "
           "buffalo_l que eu uso: são liberados só para pesquisa não comercial. Este TCC é pesquisa acadêmica, então "
           "está coberto. Quando o CESMAC colocar o sistema em operação, aí precisa licenciar — decisão da "
           "instituição, e quem confirma a leitura da licença é o jurídico. Se a licença não sair, o adaptador troca "
           "pelo YuNet + SFace (MIT e Apache 2.0), que tirou 89,5 no mesmo teste; a diferença está medida: 203 fotos "
           "erradas no topo contra 7.\n\nSe perguntarem 'o CESMAC não vende nada, então não é comercial?': a licença "
           "diz 'non-commercial RESEARCH purposes only'. A palavra que decide é pesquisa; operar um serviço para "
           "centenas de pessoas não é pesquisa, mesmo gratuito. Leitura conservadora; a palavra final é do jurídico.")

# ================================================================== 11. LGPD (new)
lgpd = new_slide("LGPD desde o desenho")
cards = [("A selfie nunca é guardada", "compara e descarta"),
         ("Sem vetor antes de publicar", "foto descartada nunca vira biometria"),
         ("Busca só dentro do evento", "nunca no acervo inteiro"),
         ("Tudo tem prazo", "expira em 6 meses · remoção em 72 h")]
for i, (t, sub) in enumerate(cards):
    x, y = LEFT + (i % 2) * 4.6, 2.0 + (i // 2) * 1.95
    box(lgpd, x, y, 4.4, 1.75, PAGE, LINE, line_w=0.75)
    circle(lgpd, x + 0.25, y + 0.5, 0.65, GREEN, "✓", size=20)
    tx(lgpd, x + 1.1, y + 0.3, 3.15, 1.2, [[(t, {"bold": True, "size": 17, "color": NAVY})],
                                         [(sub, {"size": 13, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE, spacing=4)
tx(lgpd, LEFT, 6.0, WIDTH, 0.35, "Rosto é dado pessoal sensível — LGPD, Art. 5º, II", size=12, italic=True,
   color=MUTED)
notes(lgpd, "O rosto é dado biométrico, sensível pela LGPD. Por isso: a selfie nunca é armazenada, nenhum vetor "
            "existe antes da publicação, a busca é sempre dentro de um evento, e tudo expira — com remoção a pedido em "
            "até 72 horas. Ao sair da página, o resultado também é apagado.")

# ================================================================== 12. vídeo (template slide 6)
video = T[5]
vbox = shape_named(video, "CaixaDeTexto 1")
txBody = vbox.text_frame._txBody
paras = txBody.findall(qn("a:p"))
intro_p, bullet_p = copy.deepcopy(paras[0]), copy.deepcopy(paras[1])
for p in paras:
    txBody.remove(p)
for proto_p, text in [(intro_p, "Demonstração gravada do sistema:"),
                      (bullet_p, "o fotógrafo publica um evento"),
                      (bullet_p, "o visitante encontra as próprias fotos pela selfie")]:
    p = copy.deepcopy(proto_p)
    p.findall(qn("a:r"))[0].find(qn("a:t")).text = text
    txBody.append(p)
notes(video, "Chamar o vídeo: o fotógrafo publicando um evento, e o visitante encontrando as próprias fotos pela "
             "selfie. Cuidado de LGPD: o vídeo só mostra fotos minhas ou de voluntários que concordaram.")

# ================================================================== 13. resultados parciais (template slide 8)
res = T[7]
remove(shape_named(res, "CaixaDeTexto 4"))  # "Esse slide só existe na pré-banca"


def fill_column(shape, heading, items):
    txBody = shape.text_frame._txBody
    paras = txBody.findall(qn("a:p"))
    head_p, item_p = copy.deepcopy(paras[0]), copy.deepcopy(paras[1])
    for p in paras:
        txBody.remove(p)
    for proto_p, text in [(head_p, heading)] + [(item_p, t) for t in items]:
        p = copy.deepcopy(proto_p)
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = text
        runs[0].find(qn("a:rPr")).set("sz", "1600")
        txBody.append(p)


left_col, right_col = shape_named(res, "CaixaDeTexto 2"), shape_named(res, "CaixaDeTexto 1")
for col in (left_col, right_col):
    col.top, col.width = Inches(2.35), Inches(4.5)
left_col.left, right_col.left = Inches(0.3), Inches(5.1)
fill_column(left_col, "Etapas Concluídas", ["1 · Base do sistema", "2a · Login da equipe", "2b · Área do fotógrafo",
                                            "3 · Área do visitante", "4a · Contas da equipe",
                                            "4b · Ajustes da IA (em teste)"])
fill_column(right_col, "Etapas Pendentes", ["4c · Remoções e uso atípico", "4d · Pessoas institucionais",
                                            "Plano B do reconhecimento", "Calibrar o limiar",
                                            "Aplicativo móvel", "Implantação no CESMAC"])
tx(res, LEFT, 5.85, WIDTH, 0.35, "125 testes automáticos cobrem as etapas concluídas", size=13, italic=True,
   color=MUTED, align=PP_ALIGN.CENTER)
notes(res, "Construí o sistema em etapas, na ordem em que uma depende da outra: a base, o login da equipe (2a), a "
           "área do fotógrafo (2b) e a do visitante (3). Com isso, as fases 1 e 2 do documento de requisitos estão "
           "completas. Da etapa 4, do administrador, as contas (4a) estão prontas e os temas e ajustes da IA (4b) "
           "estão feitos e em teste. Tudo isso é coberto por 125 testes automáticos. O que falta, explico no próximo "
           "slide.")

# ================================================================== 14. próximos passos (new)
nxt = new_slide("Próximos passos")
tx(nxt, LEFT, 1.8, 4.4, 0.35, "Para fechar o TCC", size=16, bold=True, color=BLUE)
tx(nxt, 5.1, 1.8, 4.4, 0.35, "Depois, na implantação", size=16, bold=True, color=MUTED)
todo = [("4c · Remoções e uso atípico", "Tela de pedidos do administrador, alerta de abuso e relatório"),
        ("4d · Pessoas institucionais", "Cadastro com autorização escrita e busca pelo nome (LGPD, Art. 11)"),
        ("Plano B do reconhecimento", "YuNet + SFace para troca planejada; se a IA falhar, só a busca pausa"),
        ("Calibrar o limiar", "Medir com mais pessoas antes de fixar o 0,30")]
later = [("Aplicativo móvel", "Mesma base em Expo; a versão web veio primeiro"),
         ("Implantação no CESMAC", "Servidor e e-mail da instituição"),
         ("Licença do modelo", "Licenciar o buffalo_l ou usar o plano B: decisão do CESMAC")]
for col_x, items, fill in [(LEFT, todo, SOFT), (5.1, later, PAGE)]:
    for i, (t, sub) in enumerate(items):
        y = 2.25 + i * 1.05
        box(nxt, col_x, y, 4.4, 0.95, fill, radius=0.1)
        tx(nxt, col_x + 0.18, y + 0.08, 4.05, 0.85, [[(t, {"bold": True, "size": 15, "color": NAVY})],
                                                   [(sub, {"size": 12, "color": "4B5966"})]], spacing=2)
notes(nxt, "Para fechar o TCC faltam quatro coisas. A 4c: o visitante já pede a remoção de uma foto, e ela some da busca "
           "na hora e é apagada em 72 horas; falta o lado do administrador — a tela dos pedidos, o alerta de uso "
           "estranho e o relatório. A 4d: um cadastro restrito de coordenadores e diretores que autorizaram por "
           "escrito, para a comunicação achar as fotos deles pelo nome (Art. 11 da LGPD). O plano B do reconhecimento: "
           "o YuNet + SFace pronto para troca planejada, e se o InsightFace falhar, só a busca pausa — o sistema nunca "
           "troca de IA sozinho, porque um modelo sem calibrar mostraria fotos de outras pessoas. E calibrar o limiar "
           "com mais gente. Depois, na implantação: aplicativo, servidor e e-mail do CESMAC, e a licença do modelo.")

# ================================================================== 15-16. referências (template slide 10 + a copy)
REFS = [
    [("BRASIL. Lei nº 13.709, de 14 de agosto de 2018. ", {}), ("Lei Geral de Proteção de Dados Pessoais (LGPD)", {"bold": True}),
     (". Brasília, DF: Presidência da República, 2018. Disponível em: https://www.planalto.gov.br/ccivil_03/_ato2015-2018/"
      "2018/lei/l13709.htm. Acesso em: 11 out. 2026.", {})],
    [("DENG, Jiankang; GUO, Jia; XUE, Niannan; ZAFEIRIOU, Stefanos. ArcFace: additive angular margin loss for deep face "
      "recognition. In: IEEE/CVF CONFERENCE ON COMPUTER VISION AND PATTERN RECOGNITION, 2019, Long Beach. ", {}),
     ("Proceedings", {"bold": True}), (" [...]. Long Beach: IEEE, 2019. p. 4690-4699.", {})],
    [("EXPO. ", {}), ("Expo Documentation", {"bold": True}),
     (". [S. l.]: 650 Industries, 2026. Disponível em: https://docs.expo.dev. Acesso em: 11 out. 2026.", {})],
    [("GUO, Jia; DENG, Jiankang; LATTAS, Alexandros; ZAFEIRIOU, Stefanos. Sample and computation redistribution for "
      "efficient face detection. In: INTERNATIONAL CONFERENCE ON LEARNING REPRESENTATIONS, 10., 2022, [S. l.]. ", {}),
     ("Proceedings", {"bold": True}), (" [...]. [S. l.]: ICLR, 2022. Disponível em: https://mlanthology.org/iclr/2022/"
                                      "guo2022iclr-sample/. Acesso em: 11 out. 2026.", {})],
    [("INSIGHTFACE. ", {}), ("InsightFace", {"bold": True}),
     (": state-of-the-art 2D and 3D face analysis project. [S. l.], 2026. Disponível em: https://github.com/deepinsight/"
      "insightface. Acesso em: 11 out. 2026.", {})],
    [("KANE, Andrew. ", {}), ("pgvector", {"bold": True}),
     (": open-source vector similarity search for Postgres. [S. l.], 2026. Disponível em: https://github.com/pgvector/"
      "pgvector. Acesso em: 11 out. 2026.", {})],
    [("MONTALA. ", {}), ("AI Faces (InsightFace)", {"bold": True}),
     (". ResourceSpace Knowledge Base. [S. l.], 2026. Disponível em: https://www.resourcespace.com/knowledge-base/"
      "plugins/faces. Acesso em: 11 out. 2026.", {})],
    [("NUNES, Thiago José Camêlo. ", {}), ("Documento de requisitos", {"bold": True}),
     (": FOTOMAC — plataforma de acervo fotográfico de eventos com busca por reconhecimento facial. Versão 1.6. "
      "Maceió: Centro Universitário CESMAC, 2026.", {})],
    [("NUNES, Thiago José Camêlo. ", {}), ("fotomac-lab", {"bold": True}),
     (": Teste 1 — comparação de tecnologias de reconhecimento facial. Maceió, 2026. Relatório técnico.", {})],
    [("OPENCV. ", {}), ("OpenCV Zoo", {"bold": True}),
     (": modelos YuNet e SFace. [S. l.], 2026. Disponível em: https://github.com/opencv/opencv_zoo. Acesso em: "
      "11 out. 2026.", {})],
    [("POSTGRESQL GLOBAL DEVELOPMENT GROUP. ", {}), ("PostgreSQL", {"bold": True}),
     (". [S. l.], 2026. Disponível em: https://www.postgresql.org. Acesso em: 11 out. 2026.", {})],
    [("RAMÍREZ, Sebastián. ", {}), ("FastAPI", {"bold": True}),
     (". [S. l.], 2026. Disponível em: https://fastapi.tiangolo.com. Acesso em: 11 out. 2026.", {})],
    [("SEAWEEDFS. ", {}), ("SeaweedFS", {"bold": True}),
     (". [S. l.], 2026. Disponível em: https://github.com/seaweedfs/seaweedfs. Acesso em: 11 out. 2026.", {})],
    [("WILL IT RUN AI. ", {}), ("Can Llama 3.2 11B Vision run on RTX 3060 12GB?", {"bold": True}),
     (" [S. l.], 2026. Disponível em: https://willitrunai.com/can-run/llama-3.2-11b-vision-on-rtx-3060-12gb. "
      "Acesso em: 11 out. 2026.", {})],
]
refs1 = T[9]
template_body = shape_named(refs1, "CustomShape 2")
remove(template_body)
refs2 = new_slide("REFERÊNCIAS", title_top=1.31)
for slide, chunk in [(refs1, REFS[:7]), (refs2, REFS[7:])]:
    tx(slide, LEFT, 2.2, WIDTH, 4.5, chunk, size=11, color="000000", spacing=9)
notes(refs1, "Referências no formato da ABNT (NBR 6023), só do que é citado nos slides.")
notes(refs2, "Continuação das referências.")

# ================================================================== 17. apêndice (new)
ap = new_slide("Apêndice: como licenciar")
ap.shapes.add_picture(str(SHOT_CONTACT), Inches(LEFT), Inches(1.95), height=Inches(3.9))
tx(ap, LEFT, 5.92, 4.6, 0.3, "insightface.ai: formulário de consulta corporativa", size=10, italic=True, color=MUTED)
blocks = [("Pedido de licença do buffalo_l", [("recognition-oss-pack@insightface.ai", {"bold": True, "color": BLUE}),
                                             ]),
          ("Preço", [("não é público: sob consulta, negociado pela instituição", {})]),
          ("Precedente", [("a Montala (ResourceSpace, sistema de acervo de imagens) comprou licença perpétua do "
                           "buffalo_l para os seus clientes", {})]),
          ("Se não licenciar", [("troca pelo YuNet + SFace e reindexa o acervo; a reindexação já existe (etapa 4b)", {})])]
for y, (t, body_runs) in zip([2.0, 3.0, 3.95, 5.15], blocks):
    tx(ap, 5.4, y, 4.1, 1.0, [[(t, {"bold": True, "size": 15, "color": NAVY})],
                              [(text, {"size": 12, "color": over.get("color", "4B5966"), "bold": over.get("bold", False)})
                               for text, over in body_runs]], spacing=2)
notes(ap, "Slide de reserva: só mostrar se perguntarem como licenciar, quanto custa ou o que acontece se a licença não "
          "sair. O print é da página de contato geral (contact@insightface.ai); para o buffalo_l, o GitHub do "
          "InsightFace indica o e-mail recognition-oss-pack@insightface.ai. Preço não é público. Referência: a Montala, "
          "do ResourceSpace (sistema de acervo de imagens, parecido com o nosso), comprou licença perpétua do buffalo_l "
          "para os clientes dela com contrato. Se não licenciar: troca pelo adaptador, reindexa e segue — perde 8 pontos "
          "no Teste 1 (97,6 para 89,5).")

# ================================================================== order, drop the template-only slides
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
# Keyed by the Python object: a slide's id can no longer be looked up once it leaves the list.
el_of = {id(s): el for s, el in zip(prs.slides, ids)}
drop = [T[6], T[8]]  # "Vídeo" PowerPoint tutorial; "Conclusão" (only at the final committee)
for s in drop:
    el = el_of[id(s)]
    prs.part.drop_rel(el.rId)
    sldIdLst.remove(el)
order = [T[0], T[1], T[2], T[3], T[4], arq, rec, t1, gen, lic, lgpd, T[5], T[7], nxt, refs1, refs2, ap]
for el in list(sldIdLst):
    sldIdLst.remove(el)
for s in order:
    sldIdLst.append(el_of[id(s)])

prs.save(str(OUT))
print(OUT, len(prs.slides), "slides")
